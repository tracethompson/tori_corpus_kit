"""
Tool 9: Export Formatter

Format analysis outputs for academic publication.
Generates MLA 8 formatted tables, figures, and quotes.
"""
import logging
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from datetime import datetime
import re
import json

import pandas as pd

try:
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

import sys
sys.path.insert(0, str(Path(__file__).parent.parent))

from utils import config
from utils.data_loader import DataLoader

logging.basicConfig(level=logging.INFO, format=config.LOG_FORMAT)
logger = logging.getLogger(__name__)


class ExportFormatter:
    """Format analysis outputs for publication."""

    def __init__(self):
        """Initialize export formatter."""
        self.data_loader = DataLoader()

        # MLA 8 citation info
        self.citation_author = config.CITATION_AUTHOR
        self.citation_title = config.CITATION_TITLE
        self.citation_docket = config.CITATION_DOCKET

    def format_mla_citation(self, comment: Dict) -> str:
        """
        Format a comment citation in MLA 8 style.

        Args:
            comment: Comment dict

        Returns:
            MLA 8 formatted citation string
        """
        submitter = comment.get("submitter_name", "Anonymous")
        if not submitter or submitter.strip() == "":
            submitter = "Anonymous"

        comment_id = comment.get("id", "")
        posted_date = comment.get("posted_date", "")

        # Format date
        if posted_date:
            try:
                dt = datetime.fromisoformat(posted_date.replace("Z", "+00:00"))
                date_str = dt.strftime("%d %b. %Y")
            except ValueError:
                date_str = posted_date
        else:
            date_str = "n.d."

        # MLA 8 format for public comment
        citation = (
            f'{submitter}. "Comment on {self.citation_title}." '
            f"Regulations.gov, {self.citation_author}, {date_str}, "
            f"www.regulations.gov/comment/{comment_id}."
        )

        return citation

    def format_quote_with_citation(
        self,
        text: str,
        comment: Dict,
        max_length: int = 300,
    ) -> Dict:
        """
        Format a quote with proper MLA 8 citation.

        Args:
            text: Quote text
            comment: Source comment
            max_length: Maximum quote length

        Returns:
            Dict with formatted quote and citation
        """
        # Truncate if needed
        if len(text) > max_length:
            # Try to break at sentence
            truncated = text[:max_length]
            last_period = truncated.rfind(".")
            if last_period > max_length * 0.5:
                text = truncated[:last_period + 1]
            else:
                text = truncated + "..."

        submitter = comment.get("submitter_name", "Anonymous")
        if not submitter or submitter.strip() == "":
            submitter = "Anonymous"

        return {
            "quote": f'"{text}"',
            "attribution": f"({submitter})",
            "full_citation": self.format_mla_citation(comment),
            "comment_id": comment.get("id", ""),
        }

    def create_table_docx(
        self,
        data: List[Dict],
        columns: List[str],
        title: str,
        output_path: Path,
        caption: Optional[str] = None,
    ) -> Path:
        """
        Create a Word document table.

        Args:
            data: List of row dicts
            columns: Column names to include
            title: Table title
            output_path: Output file path
            caption: Optional table caption

        Returns:
            Path to saved document
        """
        if not DOCX_AVAILABLE:
            logger.error("python-docx not installed")
            return None

        doc = Document()

        # Title
        title_para = doc.add_paragraph()
        title_run = title_para.add_run(title)
        title_run.bold = True
        title_run.font.size = Pt(14)
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Table
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"

        # Header row
        header_cells = table.rows[0].cells
        for i, col in enumerate(columns):
            header_cells[i].text = col
            header_cells[i].paragraphs[0].runs[0].bold = True

        # Data rows
        for row_data in data:
            row_cells = table.add_row().cells
            for i, col in enumerate(columns):
                value = row_data.get(col, "")
                row_cells[i].text = str(value)[:500]  # Truncate long values

        # Caption
        if caption:
            caption_para = doc.add_paragraph()
            caption_para.add_run(caption).italic = True

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        logger.info(f"Saved table to {output_path}")
        return output_path

    def create_table_latex(
        self,
        data: List[Dict],
        columns: List[str],
        title: str,
        output_path: Path,
        label: Optional[str] = None,
    ) -> Path:
        """
        Create a LaTeX table.

        Args:
            data: List of row dicts
            columns: Column names to include
            title: Table title
            output_path: Output file path
            label: Optional table label for references

        Returns:
            Path to saved file
        """
        # Escape LaTeX special characters
        def escape(text):
            text = str(text)
            chars = ["&", "%", "$", "#", "_", "{", "}"]
            for char in chars:
                text = text.replace(char, "\\" + char)
            return text

        # Build LaTeX
        col_spec = "|" + "l|" * len(columns)
        label_str = label or title.lower().replace(" ", "_")

        lines = [
            "\\begin{table}[htbp]",
            "\\centering",
            f"\\caption{{{escape(title)}}}",
            f"\\label{{tab:{label_str}}}",
            f"\\begin{{tabular}}{{{col_spec}}}",
            "\\hline",
            " & ".join(escape(c) for c in columns) + " \\\\",
            "\\hline",
        ]

        for row_data in data[:50]:  # Limit rows for LaTeX
            values = [escape(str(row_data.get(c, ""))[:100]) for c in columns]
            lines.append(" & ".join(values) + " \\\\")

        lines.extend([
            "\\hline",
            "\\end{tabular}",
            "\\end{table}",
        ])

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

        logger.info(f"Saved LaTeX table to {output_path}")
        return output_path

    def extract_quotes_by_theme(
        self,
        comments: List[Dict],
        themes: Dict[str, List[str]],
        max_per_theme: int = 10,
    ) -> Dict[str, List[Dict]]:
        """
        Extract representative quotes organized by theme.

        Args:
            comments: List of comment dicts
            themes: Dict mapping theme names to keyword lists
            max_per_theme: Maximum quotes per theme

        Returns:
            Dict of quotes by theme
        """
        quotes_by_theme = {theme: [] for theme in themes}

        for comment in comments:
            text = comment.get("full_text", comment.get("comment_text", ""))
            text_lower = text.lower()

            for theme, keywords in themes.items():
                if len(quotes_by_theme[theme]) >= max_per_theme:
                    continue

                # Check if comment matches theme
                matches = sum(1 for kw in keywords if kw.lower() in text_lower)
                if matches >= 2:  # Require at least 2 keyword matches
                    # Find relevant sentence
                    sentences = re.split(r"[.!?]+", text)
                    for sentence in sentences:
                        sent_lower = sentence.lower()
                        if any(kw.lower() in sent_lower for kw in keywords):
                            if 20 < len(sentence) < 300:
                                formatted = self.format_quote_with_citation(
                                    sentence.strip(), comment
                                )
                                quotes_by_theme[theme].append(formatted)
                                break

        return quotes_by_theme

    def generate_methodology_section(
        self,
        stats: Dict,
        output_path: Optional[Path] = None,
    ) -> str:
        """
        Generate methodology section text.

        Args:
            stats: Analysis statistics
            output_path: Optional output file path

        Returns:
            Methodology text
        """
        total = stats.get("total_comments", 0)
        types = stats.get("comment_types", {})
        unique_ratio = stats.get("uniqueness_ratio", 0)

        methodology = f"""
## Methodology

### Data Collection

This study analyzed {total:,} public comments submitted to FDA docket {self.citation_docket},
"Human Cells, Tissues, and Cellular and Tissue-Based Products (HCT/Ps) from Adipose Tissue:
Regulatory Considerations; Draft Guidance for Industry."

Comments were collected via the regulations.gov API between the docket opening date
and comment period close. Each comment was processed to extract full text, including
PDF attachments where applicable.

### Classification

Comments were automatically classified into three categories:
- Personal Narratives ({types.get('personal_narrative', 0):,} comments):
  First-person accounts emphasizing individual experience
- Technical Documents ({types.get('technical_document', 0):,} comments):
  Regulatory analysis with citations to CFR and PHSA
- Organizational Submissions ({types.get('organizational', 0):,} comments):
  Formal comments submitted on behalf of organizations

### Duplicate Detection

Near-duplicate comments were identified using MinHash/LSH algorithms with a
similarity threshold of {config.SIMILARITY_THRESHOLD}. This analysis found that
{unique_ratio:.1f}% of comments represented unique arguments, with the remainder
identified as variations of form letters.

### Analysis Tools

The analysis employed:
- spaCy for natural language processing (tokenization, lemmatization, NLP tagging)
- Custom frame detection for rhetorical analysis
- Multi-signal stakeholder classification

### Limitations

PDF extraction may have introduced errors in some attached documents;
these were flagged for manual review. Automated classification has inherent
limitations and a sample was manually validated.
"""

        if output_path:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(methodology)
            logger.info(f"Saved methodology to {output_path}")

        return methodology

    def generate_works_cited(
        self,
        comments: List[Dict],
        output_path: Optional[Path] = None,
    ) -> str:
        """
        Generate Works Cited entries for quoted comments.

        Args:
            comments: List of comment dicts that were quoted
            output_path: Optional output file path

        Returns:
            Works Cited text
        """
        entries = []
        for comment in comments:
            citation = self.format_mla_citation(comment)
            entries.append(citation)

        # Sort alphabetically
        entries.sort()

        works_cited = "## Works Cited\n\n" + "\n\n".join(entries)

        if output_path:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(works_cited)
            logger.info(f"Saved works cited to {output_path}")

        return works_cited

    def export_full_report(
        self,
        comments: Optional[List[Dict]] = None,
        output_dir: Optional[Path] = None,
    ) -> Dict[str, Path]:
        """
        Generate complete publication-ready export.

        Args:
            comments: List of comment dicts
            output_dir: Output directory

        Returns:
            Dict of output file paths
        """
        if comments is None:
            comments = self.data_loader.load_comments_json()

        if output_dir is None:
            output_dir = config.OUTPUTS_DIR

        config.ensure_directories()

        output_files = {}

        # Load summary stats if available
        summary_path = config.PROCESSED_DIR / "corpus_summary.json"
        if summary_path.exists():
            with open(summary_path) as f:
                stats = json.load(f)
        else:
            stats = {"total_comments": len(comments)}

        # Add uniqueness ratio if available
        dup_path = config.PROCESSED_DIR / "unique_arguments_analysis.json"
        if dup_path.exists():
            with open(dup_path) as f:
                dup_data = json.load(f)
            stats["uniqueness_ratio"] = dup_data.get("uniqueness_ratio", 0)

        # Methodology section
        path = output_dir / "methodology.md"
        self.generate_methodology_section(stats, path)
        output_files["methodology"] = path

        # Tables
        tables_dir = output_dir / "tables"

        # Comment type summary table
        type_data = [
            {"Type": k, "Count": v, "Percentage": f"{v/len(comments)*100:.1f}%"}
            for k, v in stats.get("comment_types", {}).items()
        ]
        if type_data and DOCX_AVAILABLE:
            path = tables_dir / "comment_types.docx"
            self.create_table_docx(
                type_data,
                ["Type", "Count", "Percentage"],
                "Table 1: Comment Type Distribution",
                path,
            )
            output_files["table_types_docx"] = path

            path = tables_dir / "comment_types.tex"
            self.create_table_latex(
                type_data,
                ["Type", "Count", "Percentage"],
                "Comment Type Distribution",
                path,
            )
            output_files["table_types_latex"] = path

        # Extract thematic quotes
        themes = {
            "bodily_autonomy": ["my body", "my cells", "my choice", "freedom", "autonomy"],
            "patient_access": ["access", "afford", "insurance", "cost", "available"],
            "safety_concerns": ["safe", "safety", "risk", "tested", "proven"],
            "regulatory_burden": ["burden", "expensive", "approval", "years", "impossible"],
        }

        quotes = self.extract_quotes_by_theme(comments, themes)

        # Export quotes
        quotes_dir = output_dir / "quotes"
        for theme, theme_quotes in quotes.items():
            if theme_quotes:
                path = quotes_dir / f"quotes_{theme}.json"
                self.data_loader.save_json(theme_quotes, path)
                output_files[f"quotes_{theme}"] = path

        # Works cited for quoted comments
        quoted_ids = set()
        for theme_quotes in quotes.values():
            for q in theme_quotes:
                quoted_ids.add(q.get("comment_id", ""))

        quoted_comments = [c for c in comments if c.get("id", "") in quoted_ids]
        if quoted_comments:
            path = output_dir / "works_cited.md"
            self.generate_works_cited(quoted_comments, path)
            output_files["works_cited"] = path

        logger.info(f"Export complete. Generated {len(output_files)} files.")
        return output_files


def main():
    """Main entry point for export formatter."""
    import argparse

    parser = argparse.ArgumentParser(description="Format analysis outputs for publication")
    parser.add_argument("--input", help="Input JSON file")
    parser.add_argument("--output-dir", help="Output directory")
    parser.add_argument("--cite", help="Generate citation for comment ID")
    parser.add_argument("--methodology", action="store_true", help="Generate methodology section only")

    args = parser.parse_args()

    formatter = ExportFormatter()

    input_path = Path(args.input) if args.input else None
    comments = formatter.data_loader.load_comments_json(input_path)

    if args.cite:
        comment = next((c for c in comments if c.get("id") == args.cite), None)
        if comment:
            print(formatter.format_mla_citation(comment))
        else:
            print(f"Comment {args.cite} not found")

    elif args.methodology:
        print(formatter.generate_methodology_section({"total_comments": len(comments)}))

    else:
        output_dir = Path(args.output_dir) if args.output_dir else None
        output_files = formatter.export_full_report(comments, output_dir)
        print(f"\nGenerated {len(output_files)} files:")
        for name, path in output_files.items():
            print(f"  - {name}: {path}")


if __name__ == "__main__":
    main()
